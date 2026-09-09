"""Builds sample_data/: ~90 invented invoices across six layouts and twelve
vendors, a third of them run through a scanner/phone simulation so they
have no text layer, plus a ledger export with realistic reconciliation
faults planted in it. Ground truth for every file and every fault goes to
sample_data/EXPECTED.json so tests can oracle against it.

No real company, person, document or transaction anywhere in here.
"""
from __future__ import annotations

import json
import random
import shutil
import uuid
from dataclasses import asdict, dataclass, field
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

from sample_gen.render import Renderer
from sample_gen.scanify import scanify
from sample_gen.vendors import BUYER, VENDORS, Vendor

HERE = Path(__file__).resolve().parent
OUT = HERE / "sample_data"
INV_DIR = OUT / "invoices"
WEB_SAMPLE = HERE / "web" / "public" / "sample"
WEB_SUBSET = 32  # files shipped with the browser demo; keeps "try sample data" under ~8 MB

rng = random.Random(7)

SIGNERS = ["M. Okafor", "J. Lindqvist", "R. Castellano", "A. Brennan", "S. Whitcombe"]
SITES = ["Head office", "Warehouse B", "Depot 3", "Showroom", "Annex"]
ONES = "zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen".split()
TENS = "_ _ twenty thirty forty fifty sixty seventy eighty ninety".split()


def words(n: float, currency: str) -> str:
    whole, cents = int(n), round((n - int(n)) * 100)

    def w(x: int) -> str:
        if x < 20:
            return ONES[x]
        if x < 100:
            return TENS[x // 10] + ("-" + ONES[x % 10] if x % 10 else "")
        if x < 1000:
            return ONES[x // 100] + " hundred" + (" " + w(x % 100) if x % 100 else "")
        return w(x // 1000) + " thousand" + (" " + w(x % 1000) if x % 1000 else "")

    return f"{w(whole)} {currency} and {cents:02d}/100".capitalize()


def money(fmt: str):
    def f(x: float) -> str:
        s = f"{abs(x):,.2f}"
        if fmt == "eu":
            s = s.replace(",", "X").replace(".", ",").replace("X", ".")
        elif fmt == "plain":
            s = s.replace(",", "")
        return ("-" if x < 0 else "") + s
    return f


@dataclass
class Truth:
    file: str
    vendor: str
    invoice_number: str
    invoice_date: str
    amount: float
    currency: str
    template: str
    scanned: str            # "" | "flatbed" | "phone"
    planted: list[str] = field(default_factory=list)   # fault names this file participates in


def make_lines(v: Vendor, start: date):
    n = rng.randint(6, 9) if v.template == "statement.html" else rng.randint(1, 5)
    scale = {"TRY": 40.0, "USD": 1.0, "EUR": 1.0, "GBP": 0.8, "AUD": 1.5}[v.currency]
    lines = []
    for i in range(n):
        desc = rng.choice(v.services)
        qty = rng.choice([1, 1, 1, 2, 3, 4, 6, 10, 12])
        unit = round(rng.uniform(18, 420) * scale, 2)
        if v.template == "thermal_receipt.html":
            unit = round(rng.uniform(4, 38), 2)
        lines.append({"desc": desc, "qty": qty, "unit": unit, "total": round(qty * unit, 2),
                      "date_str": (start + timedelta(days=i * 3)).strftime(v.date_format), "site": rng.choice(SITES)})
    return lines


def build_invoice(v: Vendor, counter: int, d: date, *, credit_note=False):
    lines = make_lines(v, d)
    subtotal = round(sum(l["total"] for l in lines), 2)
    tax = round(subtotal * v.tax_rate, 2)
    total = round(subtotal + tax, 2)
    if credit_note:
        for l in lines:
            l["total"], l["unit"] = -l["total"], -l["unit"]
        subtotal, tax, total = -subtotal, -tax, -total
    number = v.numbering.format(n=counter)
    ctx = dict(v=v, buyer=BUYER, number=number, date_str=d.strftime(v.date_format),
               due_str=(d + timedelta(days=rng.choice([14, 30, 30, 45]))).strftime(v.date_format),
               lines=lines, subtotal=subtotal, tax=tax, total=total, fmt=money(v.number_format), sym=v.symbol,
               credit_note=credit_note, po=(f"PO-{rng.randint(4000, 4999)}" if rng.random() < 0.4 else None),
               uuid=str(uuid.uuid5(uuid.NAMESPACE_URL, f"{v.name}/{number}")), words=words(abs(total), v.currency),
               notes=rng.choice(["Payment within 30 days.", "Thank you for your order.", "Late payments incur 2% interest.", ""]),
               signer=rng.choice(SIGNERS), time_str=f"{rng.randint(7, 18):02d}:{rng.randint(0, 59):02d}",
               card4=f"{rng.randint(1000, 9999)}", paid_stamp=None,
               brand=rng.choice(["#2b4c7e", "#3d6b4f", "#7a2e2e", "#5a4a8a", "#1f6f78", "#8a5a1f"]))
    return number, total, ctx


def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    INV_DIR.mkdir(parents=True)

    truths: list[Truth] = []
    ledger: list[list] = []          # [date, description, reference, amount, owners]  owners: files this row settles
    faults: dict[str, list] = {k: [] for k in (
        "missing_ledger_entry", "undocumented_payment", "amount_typo", "duplicate_invoice",
        "partial_payment", "combined_payment", "fx_rounding", "credit_note")}

    period_start = date(2025, 1, 6)
    counters = {v.name: rng.randint(100, 900) for v in VENDORS}
    plan = []  # (vendor, date)
    for v in VENDORS:
        k = 5 if v.template == "thermal_receipt.html" else rng.randint(6, 9)
        for _ in range(k):
            plan.append((v, period_start + timedelta(days=rng.randint(0, 170))))
    plan.sort(key=lambda t: t[1])

    def desc_for(v: Vendor, number: str) -> str:
        return rng.choice([f"Payment to {v.name}", f"{v.name.upper()} {number}", f"BACS {v.name}",
                           f"Wire transfer {number}", f"Card purchase {v.name}", f"{v.name} - invoice {number}"])

    with Renderer() as r:
        files_by_idx = []
        for idx, (v, d) in enumerate(plan):
            counters[v.name] += rng.randint(1, 4)
            number, total, ctx = build_invoice(v, counters[v.name], d)
            fname = rng.choice([f"{number.replace('/', '-')}.pdf", f"{number.replace('/', '-')}.pdf",
                                f"{v.name.split()[0]}_{d:%Y%m%d}.pdf", f"scan_{idx:04d}.pdf", f"IMG_{2000 + idx}.pdf"])
            clean = INV_DIR / fname
            r.render(v.template, ctx, clean)
            scanned = ""
            roll = rng.random()
            if roll < 0.30:
                scanned = "flatbed"
            elif roll < 0.38:
                scanned = "phone"
            if scanned:
                tmp = clean.with_suffix(".clean.pdf")
                clean.rename(tmp)
                scanify(tmp, clean, rng, mode=scanned)
                tmp.unlink()
            t = Truth(fname, v.name, number, d.isoformat(), total, v.currency, v.template, scanned)
            truths.append(t)
            files_by_idx.append((t, v, ctx))
            pay_date = d + timedelta(days=rng.randint(2, 24))
            ref = number if rng.random() < 0.8 else None
            if ref is None:
                pay_date = d + timedelta(days=rng.randint(0, 4))  # matcher's date window when no reference
            ledger.append([pay_date, desc_for(v, number), ref, total, {fname}])

        # ---- planted faults -------------------------------------------------
        eligible = [i for i, (t, v, _) in enumerate(files_by_idx) if v.template != "thermal_receipt.html"]
        rng.shuffle(eligible)
        take = lambda n: [eligible.pop() for _ in range(n)]

        for i in take(3):                                  # 1. invoice with no payment
            t = files_by_idx[i][0]
            ledger[:] = [row for row in ledger if row[2] != t.invoice_number and not (row[2] is None and row[3] == t.amount)]
            t.planted.append("missing_ledger_entry"); faults["missing_ledger_entry"].append(t.file)

        for i in take(3):                                  # 3. amount typo in the ledger
            t = files_by_idx[i][0]
            for row in ledger:
                if row[2] == t.invoice_number or (row[2] is None and row[3] == t.amount):
                    s = f"{t.amount:.2f}"
                    kind = rng.choice(["transpose", "plus50", "tenfold"])
                    if kind == "transpose" and len(s) > 4 and s[-4] != s[-5]:
                        row[3] = float(s[:-5] + s[-4] + s[-5] + s[-3:])
                    elif kind == "tenfold":
                        row[3] = round(t.amount * 10, 2)
                    else:
                        row[3] = round(t.amount + 50, 2)
                    t.planted.append(f"amount_typo:{kind}"); faults["amount_typo"].append([t.file, t.amount, row[3]])
                    break

        for i in take(2):                                  # 4. same invoice, filed twice (second copy is a scan)
            t, v, ctx = files_by_idx[i]
            dup = INV_DIR / f"{Path(t.file).stem}_rescan.pdf"
            tmp = INV_DIR / "_dup_clean.pdf"
            r.render(v.template, ctx, tmp)
            scanify(tmp, dup, rng, mode="flatbed"); tmp.unlink()
            truths.append(Truth(dup.name, v.name, t.invoice_number, t.invoice_date, t.amount, v.currency, v.template, "flatbed", ["duplicate_invoice"]))
            t.planted.append("duplicate_invoice"); faults["duplicate_invoice"].append([t.file, dup.name])

        for i in take(2):                                  # 5. paid in two instalments
            t = files_by_idx[i][0]
            for row in ledger:
                if row[2] == t.invoice_number or (row[2] is None and row[3] == t.amount):
                    first = round(t.amount * rng.choice([0.5, 0.6, 0.7]), 2)
                    row[3], row[2] = first, t.invoice_number
                    ledger.append([row[0] + timedelta(days=rng.randint(7, 20)), row[1] + " (balance)", t.invoice_number, round(t.amount - first, 2), {t.file}])
                    t.planted.append("partial_payment"); faults["partial_payment"].append([t.file, first, round(t.amount - first, 2)])
                    break

        pair = take(2)                                     # 6. two invoices, one transfer
        a, b = files_by_idx[pair[0]][0], files_by_idx[pair[1]][0]
        ledger[:] = [row for row in ledger if row[2] not in (a.invoice_number, b.invoice_number)]
        ledger.append([date.fromisoformat(max(a.invoice_date, b.invoice_date)) + timedelta(days=9),
                       f"Wire transfer {a.invoice_number} + {b.invoice_number}", f"{a.invoice_number} {b.invoice_number}",
                       round(a.amount + b.amount, 2), {a.file, b.file}])
        for t in (a, b):
            t.planted.append("combined_payment")
        faults["combined_payment"].append([a.file, b.file])

        for i in take(2):                                  # 7. FX rounding drift of a few cents
            t = files_by_idx[i][0]
            for row in ledger:
                if row[2] == t.invoice_number:
                    row[3] = round(t.amount + rng.choice([0.02, 0.03, -0.04]), 2)
                    t.planted.append("fx_rounding"); faults["fx_rounding"].append([t.file, t.amount, row[3]])
                    break

        cn_vendor = next(v for v in VENDORS if v.template == "uk_service.html")   # 8. a credit note
        counters[cn_vendor.name] += 2
        number, total, ctx = build_invoice(cn_vendor, counters[cn_vendor.name], date(2025, 5, 19), credit_note=True)
        cn_file = INV_DIR / f"{number}_CN.pdf"
        r.render(cn_vendor.template, ctx, cn_file)
        truths.append(Truth(cn_file.name, cn_vendor.name, number, "2025-05-19", total, cn_vendor.currency, cn_vendor.template, "", ["credit_note"]))
        ledger.append([date(2025, 5, 28), f"Refund {cn_vendor.name}", number, total, {cn_file.name}])
        faults["credit_note"].append(cn_file.name)

    for d_, desc, amt in [                                 # 2. payments with no invoice on file
        (date(2025, 2, 3), "Direct debit  Meridian Print Services", 312.40),
        (date(2025, 3, 28), "Bank charges", 18.00),
        (date(2025, 4, 30), "Card purchase  Saltmarsh Coffee Roasters", 46.85),
    ]:
        ledger.append([d_, desc, None, amt, set()]); faults["undocumented_payment"].append([d_.isoformat(), desc, amt])

    ledger.sort(key=lambda row: row[0])

    def write_ledger(rows: list[list], path: Path) -> None:
        wb = Workbook(); ws = wb.active; ws.title = "Transactions"
        ws.append(["Date", "Description", "Reference", "Amount"])
        for row in rows:
            ws.append(row[:4])
        for cell in ws["A"][1:]:
            cell.number_format = "yyyy-mm-dd"
        for cell in ws["D"][1:]:
            cell.number_format = "#,##0.00"
        wb.save(path)

    write_ledger(ledger, OUT / "ledger.xlsx")

    (OUT / "EXPECTED.json").write_text(json.dumps({
        "seed": 7, "invoices": [asdict(t) for t in truths], "ledger_rows": len(ledger), "faults": faults,
    }, indent=2, ensure_ascii=False))

    # browser demo subset: every planted-fault file, then fill with a spread of vendors/scans
    subset = [t for t in truths if t.planted]
    rest = [t for t in truths if not t.planted]
    rng.shuffle(rest)
    subset += rest[: max(0, WEB_SUBSET - len(subset))]
    if WEB_SAMPLE.exists():
        shutil.rmtree(WEB_SAMPLE)
    (WEB_SAMPLE / "invoices").mkdir(parents=True)
    for t in subset:
        shutil.copy(INV_DIR / t.file, WEB_SAMPLE / "invoices" / t.file)
    # the browser ledger carries only the rows that concern the subset, plus
    # every row that settles no invoice at all (the undocumented payments)
    subset_files = {t.file for t in subset}
    write_ledger([row for row in ledger if not row[4] or row[4] & subset_files], WEB_SAMPLE / "ledger.xlsx")
    (WEB_SAMPLE / "manifest.json").write_text(json.dumps({"invoices": sorted(t.file for t in subset), "ledger": "ledger.xlsx"}, indent=2))

    scanned = sum(1 for t in truths if t.scanned)
    print(f"{len(truths)} invoices ({scanned} scanned, {sum(1 for t in truths if t.scanned == 'phone')} phone photos), "
          f"{len(ledger)} ledger rows, {sum(len(v) for v in faults.values())} planted faults; web subset {len(subset)} files")


if __name__ == "__main__":
    main()
