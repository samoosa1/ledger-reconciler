"""Read a general ledger export (xlsx), auto-detecting which of the known
platform column layouts it's in — see ledger_profiles.py for the supported
platforms and how detection works.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl

from .ledger_profiles import detect_profile
from .models import LedgerRow


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _as_float(value) -> float | None:
    return float(value) if isinstance(value, (int, float)) else None


def read_ledger(xlsx_path: Path) -> list[LedgerRow]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    raw_headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [h.lower() for h in raw_headers]

    detected = detect_profile(headers)
    if detected is None:
        raise ValueError(
            "Couldn't recognize this ledger's columns against any known platform "
            f"format (Xero, Wave, FreshBooks, QuickBooks, or plain Date/Description/"
            f"Reference/Amount). Found headers: {', '.join(h for h in raw_headers if h)}"
        )
    profile, col = detected

    rows: list[LedgerRow] = []
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        values = [c.value for c in row]
        if all(v is None for v in values):
            continue

        if profile.uses_debit_credit_split:
            # We only need a matching magnitude here, not signed dr/cr
            # accounting semantics — take whichever side of the pair is
            # populated for this row.
            debit = _as_float(values[col["debit"]]) if "debit" in col else None
            credit = _as_float(values[col["credit"]]) if "credit" in col else None
            amount = debit if debit else credit
        else:
            amount = _as_float(values[col["amount"]])

        reference = values[col["reference"]] if "reference" in col else None

        rows.append(LedgerRow(
            row_index=i,
            ledger_date=_as_date(values[col["date"]]),
            description=str(values[col["description"]] or "").strip(),
            reference=str(reference).strip() if reference else None,
            amount=amount,
        ))
    return rows
